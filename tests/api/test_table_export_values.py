"""POST /tables/export over values xlsx has no cell type for: a list or a
dict is written as the text CSV writes for it."""

import csv
import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from data_rover.api.main import create_app
from data_rover.api.table_export import build_workbook
from data_rover.core.metamodel.schema import Metamodel
from data_rover.core.model.model import Model
from data_rover.core.table.cells import ValueCell

from .conftest import AUTH_HEADERS, papi, seed_default_project

_METAMODEL = """
elements:
  - name: Item
    properties:
      - {name: name, datatype: string}
      - {name: tags, datatype: string, multiplicity: "0..*"}
"""


@pytest.fixture
def client() -> TestClient:
    seed_default_project()
    c = TestClient(create_app())
    c.headers.update(AUTH_HEADERS)
    return c


def _bootstrap(client: TestClient) -> None:
    r = client.post(
        papi("/metamodel"),
        content=_METAMODEL,
        headers={"content-type": "application/x-yaml"},
    )
    assert r.status_code == 200, r.text
    client.post(papi("/model"), json={"elements": [], "relationships": []})
    r = client.post(
        papi("/model/elements"),
        json={"type": "Item", "properties": {"name": "one", "tags": ["a", "b"]}},
    )
    assert r.status_code in (200, 201), r.text


def _body(fmt: str) -> dict[str, object]:
    return {
        "definition": {
            "row_source": {"kind": "scope", "types": ["Item"]},
            "columns": [
                {"kind": "element", "header": "Item"},
                {"kind": "property", "name": "tags", "header": "Tags"},
            ],
        },
        "format": fmt,
    }


def test_xlsx_writes_a_list_value_as_its_text(client: TestClient) -> None:
    _bootstrap(client)
    r = client.post(papi("/tables/export"), json=_body("xlsx"))
    assert r.status_code == 200, r.text
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws is not None
    cell = ws["B2"]
    assert cell.value == "['a', 'b']"
    assert cell.data_type == "s"


def test_csv_writes_the_same_text_for_a_list_value(client: TestClient) -> None:
    _bootstrap(client)
    r = client.post(papi("/tables/export"), json=_body("csv"))
    assert r.status_code == 200, r.text
    rows = list(csv.reader(io.StringIO(r.content.decode("utf-8"))))
    assert rows == [["Item", "Tags"], ["one", "['a', 'b']"]]


def test_xlsx_writes_a_dict_value_as_its_text() -> None:
    blob = build_workbook(
        Model(Metamodel()), ["H"], "S", [[ValueCell(True, {"k": 1}, "e1", True)]]
    )
    ws = load_workbook(io.BytesIO(blob)).active
    assert ws is not None
    assert ws["A2"].value == "{'k': 1}"
    assert ws["A2"].data_type == "s"
