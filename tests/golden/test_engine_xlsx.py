"""The engine's committed sample workbook, opened by openpyxl, holds the grid
the frozen fixture recorded for the same export."""

from __future__ import annotations

import io
import json
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from tests.golden.reader import FIXTURE_DIR, ROOT

SAMPLE = ROOT / "engine" / "fixtures" / "xlsx" / "sample.xlsx"
CASE = "xlsx_types"


def _grid(blob: bytes) -> dict[str, Any]:
    """A one-sheet workbook as openpyxl reads it: every cell's value and type,
    the column widths it sets, the frozen pane and the autofilter range."""
    workbook = load_workbook(io.BytesIO(blob))
    (sheet,) = workbook.worksheets
    widths: dict[int, float] = {}
    for dim in sheet.column_dimensions.values():
        if dim.customWidth:
            assert dim.min is not None and dim.max is not None
            for index in range(dim.min, dim.max + 1):
                widths[index] = dim.width
    return {
        "title": sheet.title,
        "rows": [
            [{"v": cell.value, "t": cell.data_type} for cell in row]
            for row in sheet.iter_rows()
        ],
        "widths": {get_column_letter(i): widths[i] for i in sorted(widths)},
        "pane": sheet.freeze_panes,
        "autofilter": sheet.auto_filter.ref,
    }


def _recorded() -> Any:
    doc = json.loads((FIXTURE_DIR / "export_bytes.json").read_text(encoding="utf-8"))
    (step,) = [step for step in doc["steps"] if step.get("case") == CASE]
    return step["result"]["file"]["xlsx"]


def test_openpyxl_reads_the_engine_sample_as_the_fixture_recorded() -> None:
    assert SAMPLE.is_file(), "run `npm run xlsx-sample` in engine/"
    # As text, so an int read back where the oracle's was a float fails.
    assert json.dumps(_grid(SAMPLE.read_bytes())) == json.dumps(_recorded())
